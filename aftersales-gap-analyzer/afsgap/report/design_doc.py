"""Design document rendering (Markdown, optional DOCX, machine-readable JSON)."""

from __future__ import annotations

import json
import logging
from pathlib import Path

from ..models import RunResult

logger = logging.getLogger(__name__)

VERDICT_LABEL = {
    "best_in_class": "BEST IN CLASS - the current process matches or exceeds SAP standard and leading practice",
    "at_par": "AT PAR - sound, but with no advantage over SAP standard or leading practice",
    "improvement_needed": "IMPROVEMENT NEEDED - material gaps against SAP standard and/or leading practice",
    "insufficient_evidence": "INSUFFICIENT EVIDENCE - research did not return enough grounded material to judge",
}

SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}


def _table(headers: list[str], rows: list[list[str]]) -> str:
    if not rows:
        return "_None._\n"
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join([" --- "] * len(headers)) + "|"]
    for row in rows:
        cells = [(cell or "").replace("\n", " ").replace("|", "\\|") for cell in row]
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out) + "\n"


def render_markdown(result: RunResult) -> str:
    process = result.current_process
    sap = result.sap_research
    industry = result.industry_research
    analysis = result.gap_analysis
    lines: list[str] = []
    add = lines.append

    add(f"# Process Design & Gap Analysis - {result.process_name}")
    add("")
    add(f"**Process ID:** {result.process_id}  ")
    add(f"**Business area:** {process.business_area}  ")
    add(f"**Transition:** {process.source_system} to {process.target_system}  ")
    add(f"**Run date:** {result.run_date.isoformat()}  ")
    add(f"**Mode:** {'offline (fixtures)' if result.offline else 'live research'}")
    add("")
    add("<!-- afsgap:meta -->")
    add("> **Scope exclusions applied to this document:** tolerance topics are excluded "
        "end to end (input, research, analysis and output). Industry benchmark content is "
        "qualitative and carries no KPIs or metrics. Transaction codes appear only where a "
        "public SAP source literally contains them.")
    add("<!-- /afsgap:meta -->")
    add("")

    # 1 -------------------------------------------------------------------
    add("## 1. Verdict")
    add("")
    add(f"**{VERDICT_LABEL.get(analysis.verdict, analysis.verdict)}**")
    add("")
    add(analysis.verdict_rationale)
    add("")
    add("### Executive summary")
    add("")
    add(analysis.executive_summary)
    add("")

    # 2 -------------------------------------------------------------------
    add("## 2. Current (AS-IS) process")
    add("")
    if process.description:
        add(process.description)
        add("")
    if process.scope_in or process.scope_out:
        add(f"**In scope:** {', '.join(process.scope_in) or '-'}  ")
        add(f"**Out of scope:** {', '.join(process.scope_out) or '-'}")
        add("")
    add(_table(
        ["#", "Step", "Actor", "System", "Description", "Pain points"],
        [
            [str(step.seq), step.name, step.actor, step.system, step.description, "; ".join(step.pain_points)]
            for step in process.steps
        ],
    ))
    if process.known_pain_points:
        add("**Known pain points**")
        add("")
        for item in process.known_pain_points:
            add(f"- {item}")
        add("")
    if process.custom_objects:
        add("**Custom objects in ECC today**")
        add("")
        for item in process.custom_objects:
            add(f"- {item}")
        add("")

    # 3 -------------------------------------------------------------------
    add("## 3. SAP standard reference")
    add("")
    add("_Sourced exclusively from official public SAP domains. Area names such as SAP SD, "
        "SAP EWM, SAP CMH, master data, the dealer front-end/Sales Order Plus and aftersales "
        "were used as search guidance only; only findings actually returned by a public SAP "
        "source appear below._")
    add("")
    if sap.standard.summary:
        add(sap.standard.summary)
        add("")
    add("### 3.1 Standard process flow")
    add("")
    add(_table(
        ["#", "Step", "SAP area", "Description", "Source"],
        [
            [str(step.seq), step.name, step.sap_area, step.description,
             ", ".join(e.url for e in step.evidence)]
            for step in sap.standard.standard_process_flow
        ],
    ))
    add("### 3.2 Relevant SAP modules and components")
    add("")
    add(_table(
        ["Module / component", "Role in this process", "Source"],
        [[m.name, m.role_in_process, ", ".join(e.url for e in m.evidence)] for m in sap.standard.modules],
    ))
    add("### 3.3 Transaction codes (verified)")
    add("")
    add("_Each code below was found literally on the cited public SAP page. Codes that could not "
        "be verified are listed underneath with the reason, and are not used anywhere else in "
        "this document._")
    add("")
    add(_table(
        ["T-code", "Purpose", "SAP area", "Evidence"],
        [[c.tcode, c.purpose, c.sap_area, c.source_url] for c in sap.verified_tcodes],
    ))
    if sap.dropped_tcodes:
        add("**Rejected transaction codes**")
        add("")
        add(_table(["Candidate", "Rejected because", "Cited source"],
                   [[d.tcode, d.reason, d.source_url] for d in sap.dropped_tcodes]))
    if sap.standard.master_data_objects:
        add("### 3.4 Master data objects")
        add("")
        for item in sap.standard.master_data_objects:
            add(f"- {item}")
        add("")
    if sap.standard.integration_points:
        add("### 3.5 Integration points")
        add("")
        for item in sap.standard.integration_points:
            add(f"- {item}")
        add("")
    if sap.standard.s4_specific_changes:
        add("### 3.6 What changes in S/4HANA versus ECC")
        add("")
        for item in sap.standard.s4_specific_changes:
            add(f"- {item}")
        add("")
    if sap.standard.open_questions:
        add("### 3.7 Open questions for SAP / the functional team")
        add("")
        for item in sap.standard.open_questions:
            add(f"- {item}")
        add("")

    # 4 -------------------------------------------------------------------
    add("## 4. Industry practice benchmark")
    add("")
    add("_Qualitative only. Sources were found through several broad search framings and then "
        "filtered against a credible-source allowlist (consultancies, standards and industry "
        "bodies, research institutes, recognised industry press), supplemented by peer-reviewed "
        "work retrieved from OpenAlex._")
    add("")
    if industry.benchmark.summary:
        add(industry.benchmark.summary)
        add("")
    add("### 4.1 Leading practices")
    add("")
    add(_table(
        ["Practice", "Maturity", "Description", "Source"],
        [
            [p.title, p.maturity, p.description, ", ".join(e.url for e in p.evidence)]
            for p in industry.benchmark.leading_practices
        ],
    ))
    for heading, items in (
        ("4.2 Common operating models", industry.benchmark.common_operating_models),
        ("4.3 Automation and digital enablers", industry.benchmark.automation_and_digital_enablers),
        ("4.4 Recurring risks and failure modes", industry.benchmark.risks_and_failure_modes),
    ):
        if items:
            add(f"### {heading}")
            add("")
            for item in items:
                add(f"- {item}")
            add("")
    if industry.benchmark.scholarly_works:
        add("### 4.5 Peer-reviewed grounding (OpenAlex)")
        add("")
        add(_table(
            ["Title", "Year", "Venue", "DOI / URL"],
            [[w.title, str(w.year or ""), w.venue, w.doi or w.url] for w in industry.benchmark.scholarly_works],
        ))

    # 5 -------------------------------------------------------------------
    add("## 5. Gap analysis")
    add("")
    gaps = sorted(analysis.gaps, key=lambda g: SEVERITY_ORDER.get(g.severity, 9))
    add(_table(
        ["ID", "Dimension", "Gap", "Severity", "S/4HANA disposition", "Effort"],
        [[g.gap_id, g.dimension, g.title, g.severity, g.s4_disposition, g.effort] for g in gaps],
    ))
    for gap in gaps:
        add(f"### {gap.gap_id} - {gap.title}")
        add("")
        add(f"**Dimension:** {gap.dimension} | **Severity:** {gap.severity} | "
            f"**Disposition:** {gap.s4_disposition} | **Effort:** {gap.effort}")
        add("")
        add(f"**Today:** {gap.current_state}")
        add("")
        add(f"**SAP standard:** {gap.sap_standard_reference}")
        add("")
        if gap.industry_reference:
            add(f"**Industry practice:** {gap.industry_reference}")
            add("")
        add(f"**Gap:** {gap.gap_description}")
        add("")
        add(f"**Recommendation:** {gap.recommendation}")
        add("")
        if gap.dependencies:
            add(f"**Dependencies:** {', '.join(gap.dependencies)}")
            add("")
        if gap.evidence:
            add("**Evidence**")
            add("")
            for evidence in gap.evidence:
                add(f"- \"{evidence.quote}\" - {evidence.url}")
            add("")

    # 6 -------------------------------------------------------------------
    add("## 6. TO-BE process design")
    add("")
    add(_table(
        ["#", "Step", "Actor", "System", "Description", "Change vs today"],
        [
            [str(s.seq), s.name, s.actor, s.system, s.description, s.change_vs_current]
            for s in analysis.to_be_process
        ],
    ))

    # 7 -------------------------------------------------------------------
    add("## 7. Recommendations and roadmap")
    add("")
    if analysis.quick_wins:
        add("### Quick wins")
        add("")
        for item in analysis.quick_wins:
            add(f"- {item}")
        add("")
    if analysis.roadmap_phases:
        add("### Phasing")
        add("")
        for index, item in enumerate(analysis.roadmap_phases, start=1):
            add(f"{index}. {item}")
        add("")
    if analysis.risks:
        add("### Risks")
        add("")
        for item in analysis.risks:
            add(f"- {item}")
        add("")
    if analysis.assumptions:
        add("### Assumptions")
        add("")
        for item in analysis.assumptions:
            add(f"- {item}")
        add("")

    # 8 -------------------------------------------------------------------
    add("## 8. Evidence register")
    add("")
    add("### 8.1 SAP sources used")
    add("")
    add(_table(["Domain", "Title", "URL"], [[s.domain, s.title, s.url] for s in sap.sources]))
    add("### 8.2 Industry sources retained")
    add("")
    add(_table(["Category", "Domain", "Title", "URL"],
               [[s.category, s.domain, s.title, s.url] for s in industry.sources]))
    if industry.rejected_sources:
        add("### 8.3 Sources rejected by the credibility filter")
        add("")
        add(_table(["Domain", "Tier", "URL"],
                   [[s.domain, s.tier, s.url] for s in industry.rejected_sources]))
    add("### 8.4 Queries run")
    add("")
    for query in sap.queries_used:
        add(f"- [SAP] {query}")
    for query in industry.queries_used:
        add(f"- [Industry] {query}")
    add("")

    # 9 -------------------------------------------------------------------
    add("## 9. Exclusion and validation log")
    add("")
    validation = result.validation
    add(f"**Validation status:** {'PASSED' if validation.passed else 'FAILED'}")
    add("")
    if validation.errors:
        add("**Errors**")
        add("")
        for item in validation.errors:
            add(f"- {item}")
        add("")
    if validation.warnings:
        add("**Warnings**")
        add("")
        for item in validation.warnings:
            add(f"- {item}")
        add("")
    add("**Content excluded by rule**")
    add("")
    add(_table(
        ["Rule", "Stage", "Location", "Excluded content", "Matched"],
        [
            [e.rule, e.stage, e.location, (e.excluded_text[:160] + ("..." if len(e.excluded_text) > 160 else "")),
             ", ".join(e.matched_terms)]
            for e in validation.exclusions
        ],
    ))
    add("")
    add("---")
    add("")
    add("_Generated by afsgap. Every SAP statement is traceable to a public SAP source; every "
        "industry statement to a credible publisher or peer-reviewed work. Unverifiable claims "
        "were dropped rather than softened._")
    return "\n".join(lines) + "\n"


def write_reports(result: RunResult, output_dir: Path, base_name: str | None = None) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = base_name or f"{result.process_id}_{result.run_date.isoformat()}"
    written: dict[str, Path] = {}

    markdown_path = output_dir / f"{stem}_design_document.md"
    markdown_path.write_text(render_markdown(result), encoding="utf-8")
    written["markdown"] = markdown_path

    json_path = output_dir / f"{stem}_run.json"
    json_path.write_text(json.dumps(result.model_dump(mode="json"), indent=2, ensure_ascii=False), encoding="utf-8")
    written["json"] = json_path

    docx_path = _write_docx(result, output_dir / f"{stem}_design_document.docx")
    if docx_path:
        written["docx"] = docx_path
    return written


def _write_docx(result: RunResult, path: Path) -> Path | None:
    try:
        from docx import Document
    except ImportError:
        logger.info("python-docx not installed - skipping DOCX export.")
        return None

    document = Document()
    for line in render_markdown(result).splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            document.add_heading(stripped.lstrip("# ").strip(), level=min(level, 4))
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("|"):
            document.add_paragraph(stripped)
        elif stripped:
            document.add_paragraph(stripped)
    document.save(path)
    return path
